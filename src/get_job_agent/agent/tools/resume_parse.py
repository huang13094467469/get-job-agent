"""简历文档解析：附件 PDF/DOCX → 纯文本。"""

from __future__ import annotations

from pathlib import Path

from langchain_core.tools import tool

SUPPORTED_SUFFIXES = {".pdf": "PDF", ".docx": "DOCX"}


def _parse_pdf(path: str) -> str:
    import fitz  # pymupdf

    doc = fitz.open(path)
    pages = [page.get_text() for page in doc]
    doc.close()
    return "\n".join(p for p in pages if p)


def _parse_docx(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text:
            parts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)
    return "\n".join(parts)


@tool
def parse_resume(file_path: str) -> str:
    """把本地简历附件(PDF/DOCX)解析为纯文本。

    参数 file_path 为附件下载到本地的绝对路径。
    返回抽取出的纯文本；格式不支持或读取失败时抛出清晰异常。
    """
    p = Path(file_path)
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"不支持的简历格式: {suffix}，仅支持 PDF/DOCX")
    if not p.exists():
        raise FileNotFoundError(f"附件不存在: {file_path}")

    if suffix == ".pdf":
        text = _parse_pdf(str(p))
    else:
        text = _parse_docx(str(p))

    text = text.strip()
    if not text:
        raise ValueError("未能从附件中抽取到文本，可能是扫描件（当前不支持 OCR）")
    return text


__all__ = ["parse_resume", "SUPPORTED_SUFFIXES"]